"""Testele exportului Excel stilizat: scriitorul (xlsx-stil.js) și meniul ⚙ de pe panourile dashboardului.

    python teste_xlsx.py             # scriitorul .xlsx, rulat cu Node: structură, formule, grafice, caractere speciale
    python teste_xlsx.py --browser   # plus: panourile din dashboard, descărcate prin ⚙ → „Descarcă Excel", în română și engleză

Are nevoie de Node.js și de openpyxl (iar pentru --browser, de Playwright). Fișierele sunt verificate cu zipfile, cu un
parser XML și cu openpyxl; potrivirea cu schema ISO 29500 și randarea în LibreOffice se fac manual, nu aici.
"""
import functools, json, os, re, shutil, subprocess, sys, tempfile, threading, time, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

AICI = Path(__file__).resolve().parent
picat = []


def verifica(nume, conditie, detaliu=''):
    print(('  ok   ' if conditie else '  PICĂ ') + nume + ('' if conditie else f'  → {detaliu}'))
    if not conditie:
        picat.append(nume)


NS_C = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}


def parti_xml_valide(cale):
    """Fiecare parte XML a arhivei trebuie să fie bine formată; întoarce lista celor care nu sunt."""
    rele = []
    with zipfile.ZipFile(cale) as z:
        for n in z.namelist():
            if n.endswith('.xml') or n.endswith('.rels'):
                try:
                    ET.fromstring(z.read(n))
                except ET.ParseError as e:
                    rele.append((n, str(e)[:60]))
    return rele


def formule_grafic(cale):
    """Adresele (c:f) din chart1.xml, pe tipuri: categorii, valori."""
    with zipfile.ZipFile(cale) as z:
        if 'xl/charts/chart1.xml' not in z.namelist():
            return None
        rad = ET.fromstring(z.read('xl/charts/chart1.xml'))
    cat = [e.text for e in rad.findall('.//c:ser/c:cat//c:f', NS_C)]
    val = [e.text for e in rad.findall('.//c:ser/c:val//c:f', NS_C)]
    tip = [t for t in ('barChart', 'lineChart') if rad.find('.//c:' + t, NS_C) is not None]
    directie = rad.find('.//c:barDir', NS_C)
    grupare = rad.find('.//c:grouping', NS_C)
    return {'cat': cat, 'val': val, 'tip': tip[0] if tip else None, 'dir': directie.get('val') if directie is not None else None,
            'grup': grupare.get('val') if grupare is not None else None}


# ============================================================ 1. scriitorul, cu Node
print('Scriitorul .xlsx (xlsx-stil.js), cu Node')
if not shutil.which('node'):
    print('  PICĂ nu există Node.js'); sys.exit(1)

GEN = r'''
const X = require(process.argv[2]);
const fs = require('fs');
const out = process.argv[3];
const salveaza = async (nume, spec) => { const b = await X.construieste(spec); fs.writeFileSync(out + '/' + nume + '.xlsx', Buffer.from(await b.arrayBuffer())); return b.size; };
(async () => {
  const donori = ['Banca Europeană de Investiții', 'Banca Mondială', 'Uniunea Europeană', 'USAID', 'GIZ', 'Sida', 'PNUD', 'JICA', 'SDC', 'Regatul Unit', 'ONU', 'BERD', 'OSCE', 'Franța'];
  await salveaza('top', {
    limba: 'ro', foaie: 'Top donatori', titlu: 'Top donatori · sumă totală', subtitlu: 'Sursa: AMP · sume în EUR', antete: ['Donator', 'Angajamente (EUR)', 'Debursări (EUR)', 'Proiecte'],
    tipuri: 'teen', randuri: donori.map((d, i) => [d, 1000000 * (20 - i), 500000 * (20 - i), 10 + i]).concat([['Total (proiecte unice)', 17000000, 8000000, 120]]), randTotal: true,
    bare: [2], note: ['O notă obișnuită.', { t: 'Deschide dashboardul', u: 'https://nistor.vivi.md/' }],
    linkuri: { 0: (r, i) => i % 2 ? null : 'https://example.org/d/' + i },
    grafic: { tip: 'bar', cat: 0, serii: [1, 2], n: 12, culori: ['9DB8DF', '1E4B8C'], titlu: 'Top 12 donatori' }
  });
  const ani = []; for (let a = 1993; a <= 2026; a++) ani.push([a, a * 1000, a * 500]);
  await salveaza('linie', {
    limba: 'ro', foaie: 'Evoluție', titlu: 'Evoluție în timp', subtitlu: 'x', antete: ['An', 'Angajamente', 'Debursări'], tipuri: 'yee', randuri: ani,
    total: { eticheta: 'Total', suma: [1, 2] }, grafic: { tip: 'line', cat: 0, serii: [1, 2], ultimele: 14 }
  });
  await salveaza('coloane', {
    limba: 'ro', foaie: 'Tip pe ani', titlu: 'Tip pe ani', subtitlu: 'x', antete: ['An', 'A', 'B', 'C', 'Total'], tipuri: 'yeeee',
    randuri: ani.map(r => [r[0], 10, 20, 30, 60]), total: { eticheta: 'Total', suma: [1, 2, 3, 4] }, grafic: { tip: 'colst', cat: 0, serii: [1, 2, 3] }
  });
  await salveaza('matrice', {
    limba: 'en', foaie: 'Heatmap', titlu: 'Donor × year', subtitlu: 'x', antete: ['Donor', '2020', '2021', '2022', 'Total'], tipuri: 'teeee',
    randuri: donori.map((d, i) => [d, i * 10, i * 20, i * 30, i * 60]), harta: { de: 1, pana: 3 }
  });
  await salveaza('semn', {
    limba: 'ro', foaie: 'Semn', titlu: 'Semn', subtitlu: 'x', antete: ['Nume', 'A', 'B', 'Dif x', 'Dif s', 'Var v', 'Loc k', 'Cu 2 zecimale', 'Procent'], tipuri: 'teexsvkdp',
    randuri: [['a', 100, 50, -50, -50, -50, 1, 1.456, 83.5], ['b', 10, 30, 20, 20, 200, 2, 2.5, 12]], randTotal: false
  });
  await salveaza('caractere', {
    limba: 'ro', foaie: 'A/B:C*D?[E]\\F oarte lung peste treizeci și unu de caractere', titlu: 'Titlu cu <tag> & "ghilimele"', subtitlu: 'Sub & titlu', subsol: '127.0.0.1:8000 & co',
    antete: ['Text', 'Număr'], tipuri: 'tn',
    randuri: [['=1+1', 1], ['+cmd|calc', 2], ['@SUM(A1)', 3], ['-5+3', 4], ['<script>alert("x")</script> & \'apostrof\'', 5], ['rând 1\nrând 2', 6], ['emoji 🎉 și diacritice ăâîșț', 7],
              ['control\u0001\u0002\u0008 scos', 8], ['x'.repeat(40000), 9], ['spații   multiple', 10], ['', 11], [null, 12]]
  });
  const mari = []; for (let i = 0; i < 20000; i++) mari.push(['idea_' + i, 'Proiect ' + i + ' — titlu', 'Donator ' + (i % 50), String(2000 + i % 25), i * 3, i * 2]);
  const t0 = Date.now();
  const dim = await salveaza('mare', {
    limba: 'ro', foaie: 'Mare', titlu: 'Lista mare', subtitlu: 'x', antete: ['AMP ID', 'Proiect', 'Donator', 'Perioadă', 'Ang', 'Deb'], tipuri: 'itttee', randuri: mari,
    total: { eticheta: 'Total', suma: [4, 5] }, linkuri: { 0: r => 'https://amp.gov.md/' + r[0] }
  });
  console.log(JSON.stringify({ mare_ms: Date.now() - t0, mare_bytes: dim }));
  await salveaza('gol', { limba: 'ro', foaie: 'Gol', titlu: 'Fără rânduri', subtitlu: 'x', antete: ['A', 'B'], tipuri: 'tn', randuri: [] });
})().catch(e => { console.error('EROARE', e && e.stack || e); process.exit(1); });
'''
tmp = Path(tempfile.mkdtemp(prefix='teste_xlsx_'))
try:
    (tmp / 'gen.js').write_text(GEN, encoding='utf-8')
    r = subprocess.run(['node', str(tmp / 'gen.js'), str(AICI / 'xlsx-stil.js'), str(tmp)], capture_output=True, text=True, timeout=180)
    verifica('Node generează toate fișierele de probă', r.returncode == 0, (r.stderr or r.stdout)[-300:])
    timp = json.loads(r.stdout.strip().splitlines()[-1]) if r.returncode == 0 else {}

    for nume in ('top', 'linie', 'coloane', 'matrice', 'semn', 'caractere', 'mare', 'gol'):
        cale = tmp / (nume + '.xlsx')
        ok_zip = cale.exists() and zipfile.is_zipfile(cale)
        verifica(f'{nume}: arhivă ZIP validă, cu toate părțile XML bine formate', ok_zip and not parti_xml_valide(cale), parti_xml_valide(cale) if ok_zip else 'lipsă')

    # ---- top: total dat gata, bare, grafic orizontal, linkuri, note
    wb = openpyxl.load_workbook(tmp / 'top.xlsx'); ws = wb.worksheets[0]
    verifica('top: titlul în A1, subtitlul în A2, antetul pe rândul 4', ws['A1'].value == 'Top donatori · sumă totală' and 'Sursa: AMP' in ws['A2'].value and ws['B4'].value == 'Angajamente (EUR)')
    verifica('top: antetul rămâne la derulare (pane înghețat) și are filtre', ws.freeze_panes == 'A5' and ws.auto_filter.ref == 'A4:D18', (ws.freeze_panes, ws.auto_filter.ref))
    verifica('top: sumele au separator de mii, nu zecimale', ws['B5'].number_format == '#,##0' and ws['B5'].value == 20000000, (ws['B5'].number_format, ws['B5'].value))
    verifica('top: rândul de total e scris ca atare, în afara filtrului', ws['A19'].value == 'Total (proiecte unice)' and ws['B19'].value == 17000000, (ws['A19'].value, ws['B19'].value))
    legaturi = [c.coordinate for row in ws.iter_rows() for c in row if c.hyperlink]
    verifica('top: hyperlink doar pe rândurile cerute (+ nota cu adresă)', legaturi[:3] == ['A5', 'A7', 'A9'] and legaturi[-1].startswith('A2'), legaturi)
    g = formule_grafic(tmp / 'top.xlsx')
    verifica('top: grafic cu bare orizontale, 2 serii, primele 12 rânduri', g and g['tip'] == 'barChart' and g['dir'] == 'bar' and len(g['val']) == 2 and all(re.search(r'\$[A-D]\$5:\$[A-D]\$16$', v) for v in g['val']), g)
    verifica('top: graficul nu include rândul de total', g and not any('$17' in v or '$18' in v for v in g['val']), g)

    # ---- linie: totaluri cu formule, grafic în linie pe ultimii 14 ani
    wb = openpyxl.load_workbook(tmp / 'linie.xlsx'); ws = wb.worksheets[0]
    verifica('linie: totalul e formulă SUBTOTAL (se recalculează la filtrare)', ws['B39'].value == '=SUBTOTAL(109,B5:B38)' and ws['C39'].value == '=SUBTOTAL(109,C5:C38)', (ws['B39'].value, ws['C39'].value))
    verifica('linie: anul e număr întreg fără separator de mii', ws['A5'].value == 1993 and ws['A5'].number_format == '0', (ws['A5'].value, ws['A5'].number_format))
    wb2 = openpyxl.load_workbook(tmp / 'linie.xlsx', data_only=True)
    verifica('linie: valoarea calculată din fișier e suma corectă (se vede și fără recalculare)', wb2.worksheets[0]['B39'].value == sum(a * 1000 for a in range(1993, 2027)), wb2.worksheets[0]['B39'].value)
    g = formule_grafic(tmp / 'linie.xlsx')
    verifica('linie: grafic în linie, ultimii 14 ani', g and g['tip'] == 'lineChart' and all(re.search(r'\$[A-D]\$25:\$[A-D]\$38$', v) for v in g['val']), g)

    # ---- coloane stivuite
    g = formule_grafic(tmp / 'coloane.xlsx')
    verifica('coloane: grafic cu coloane stivuite și 3 serii', g and g['tip'] == 'barChart' and g['dir'] == 'col' and g['grup'] == 'stacked' and len(g['val']) == 3, g)

    # ---- matrice: scară de culoare
    ws = openpyxl.load_workbook(tmp / 'matrice.xlsx').worksheets[0]
    reguli = [r.type for rng in ws.conditional_formatting for r in rng.rules]
    verifica('matrice: scară de culoare pe blocul de coloane + zebră', 'colorScale' in reguli and 'expression' in reguli, reguli)
    with zipfile.ZipFile(tmp / 'matrice.xlsx') as z:
        verifica('matrice (EN): subsolul scrie „Page”', 'Page' in z.read('xl/worksheets/sheet1.xml').decode('utf-8'))

    # ---- semn: formate și reguli verde/roșu doar pe x și v
    ws = openpyxl.load_workbook(tmp / 'semn.xlsx').worksheets[0]
    verifica('semn: formatele diferențelor și ale procentelor', ws['D5'].number_format == '+#,##0;-#,##0;0' and ws['E5'].number_format == '+#,##0;-#,##0;0'
             and ws['F5'].number_format == '+0.0%;-0.0%;0.0%' and ws['H5'].number_format == '#,##0.00' and ws['I5'].number_format == '0.0%',
             [ws[c + '5'].number_format for c in 'DEFHI'])
    verifica('semn: procentele se pun ca fracții (83,5 → 0,835)', abs(ws['I5'].value - 0.835) < 1e-9 and abs(ws['F5'].value + 0.5) < 1e-9, (ws['I5'].value, ws['F5'].value))
    celule_cf = sorted({str(rng.sqref) for rng in ws.conditional_formatting for r in rng.rules if r.type == 'cellIs'})
    verifica('semn: verde/roșu doar pe tipurile cu semn „x” și „v”, nu și pe „s”', celule_cf == ['D5:D6', 'F5:F6'], celule_cf)

    # ---- caractere speciale
    ws = openpyxl.load_workbook(tmp / 'caractere.xlsx').worksheets[0]
    verifica('caractere: numele foii fără caractere interzise și cel mult 31 de caractere', len(ws.title) <= 31 and not re.search(r'[\\/*?:\[\]]', ws.title), ws.title)
    val = [ws.cell(row=r, column=1).value for r in range(5, 17)]
    verifica('caractere: textul care începe cu = + @ - rămâne text, nu formulă', all(ws.cell(row=5 + i, column=1).data_type != 'f' for i in range(4)), [ws.cell(row=5 + i, column=1).data_type for i in range(4)])
    verifica('caractere: <, &, ghilimele, apostrof, rând nou, emoji și diacritice se păstrează', "<script>alert(\"x\")</script> & 'apostrof'" in val and 'rând 1\nrând 2' in val and 'emoji 🎉 și diacritice ăâîșț' in val, val[3:8])
    verifica('caractere: caracterele de control sunt scoase (altfel Excel respinge fișierul)', any(v and v.startswith('control') and '\u0001' not in v for v in val), val[7])
    verifica('caractere: un text de 40.000 de caractere e tăiat la limita Excel (32.767)', any(v and v.startswith('xxxx') and set(v.rstrip('…')) == {'x'} and len(v) <= 32767 and v.endswith('…') for v in val), [len(v) for v in val if v and v.startswith('xxxx')])
    verifica('caractere: titlul cu <tag> și ghilimele ajunge întreg', ws['A1'].value == 'Titlu cu <tag> & "ghilimele"', ws['A1'].value)
    with zipfile.ZipFile(tmp / 'caractere.xlsx') as z:
        subsol = re.search(r'<oddFooter>(.*?)</oddFooter>', z.read('xl/worksheets/sheet1.xml').decode('utf-8')).group(1)
    # „&8” urmat de o cifră ar fi citit ca dimensiune de font (8127…), iar „&” singur începe un cod: textul cere spațiu și „&&”
    verifica('caractere: subsolul care începe cu cifre nu se amestecă cu dimensiunea fontului, iar „&” e dublat', '&amp;L&amp;8 127.0.0.1:8000 &amp;&amp; co' in subsol, subsol)

    # ---- mare
    t = timp.get('mare_ms', 1e9)
    verifica('mare: 20.000 de rânduri cu hyperlink pe fiecare, generate în sub 20 s (%.1f s)' % (t / 1000), t < 20000, t)
    ws = openpyxl.load_workbook(tmp / 'mare.xlsx', read_only=True).worksheets[0]
    n = sum(1 for _ in ws.iter_rows(min_row=5, max_col=1, values_only=True))
    verifica('mare: toate rândurile sunt în fișier (20.000 + total)', n == 20001, n)

    # ---- gol
    ws = openpyxl.load_workbook(tmp / 'gol.xlsx').worksheets[0]
    verifica('gol: și fără rânduri iese un fișier valid, cu antetul', ws['A4'].value == 'A' and ws['B4'].value == 'B')

    # ---- un fișier cu mai multe foi
    GEN_FOI = r"""
const X = require(process.argv[2]);
const fs = require('fs');
(async () => {
  const sect = ['Educație', 'Sănătate și protecție socială (extins)', 'Agricultură/Dezvoltare rurală', 'educație'];
  const foiSector = sect.map((nm, i) => ({
    foaie: nm, titlu: nm, subtitlu: 'Sector ' + (i + 1), antete: ['AMP ID', 'Proiect', 'Sumă (EUR)'], tipuri: 'itn', culoareFila: i % 2 ? 'FFC8871A' : 'FF1E4B8C',
    randuri: [['A' + i, 'Proiect ' + i + '.1', 100 * (i + 1)], ['B' + i, 'Proiect ' + i + '.2', 200 * (i + 1)]],
    total: { eticheta: 'Total', suma: [2] }, linkuri: { 0: r => 'https://amp.gov.md/' + r[0] },
    grafic: i === 1 ? { tip: 'col', cat: 0, serii: [2] } : null
  }));
  const b = await X.construieste({
    limba: 'ro', autor: 'Test', subsol: 'nistor.vivi.md', titlu: 'Fișier cu mai multe foi',
    foi: [
      { foaie: 'Rezumat', titlu: 'Rezumat', subtitlu: 'x', antete: ['Sector', 'Proiecte', 'Sumă (EUR)'], tipuri: 'tin',
        randuri: sect.map((nm, i) => [nm, 2, 300 * (i + 1)]).concat([['Fără foaie', 1, 5]]),
        linkuri: { 0: r => r[0] === 'Fără foaie' ? '#Nu există' : '#' + r[0] },
        total: { eticheta: 'Total', suma: [1, 2] }, grafic: { tip: 'bar', cat: 0, serii: [2], titlu: 'Pe sectoare' } },
      { foaie: 'Toate proiectele', titlu: 'Toate', subtitlu: 'x', antete: ['An', 'Sumă'], tipuri: 'yn', randuri: [[2020, 5], [2021, 7], [2022, 9]],
        total: { eticheta: 'Total', suma: [1] }, grafic: { tip: 'line', cat: 0, serii: [1] }, culoareFila: 'FF2E7D32' }
    ].concat(foiSector)
  });
  fs.writeFileSync(process.argv[3] + '/foi.xlsx', Buffer.from(await b.arrayBuffer()));
})().catch(e => { console.error('EROARE', e && e.stack || e); process.exit(1); });
"""
    (tmp / 'gen_foi.js').write_text(GEN_FOI, encoding='utf-8')
    r = subprocess.run(['node', str(tmp / 'gen_foi.js'), str(AICI / 'xlsx-stil.js'), str(tmp)], capture_output=True, text=True, timeout=120)
    verifica('foi: Node generează fișierul cu 6 foi', r.returncode == 0, (r.stderr or r.stdout)[-300:])
    cale = tmp / 'foi.xlsx'
    verifica('foi: arhivă ZIP validă, cu toate părțile XML bine formate', cale.exists() and zipfile.is_zipfile(cale) and not parti_xml_valide(cale), parti_xml_valide(cale) if cale.exists() else 'lipsă')
    wb = openpyxl.load_workbook(cale)
    nume = wb.sheetnames
    verifica('foi: toate cele 6 foi, în ordine', len(nume) == 6 and nume[0] == 'Rezumat' and nume[1] == 'Toate proiectele', nume)
    verifica('foi: numele sunt unice (fără deosebire între litere mari și mici) și de cel mult 31 de caractere',
             len({n.lower() for n in nume}) == len(nume) and all(len(n) <= 31 for n in nume), nume)
    verifica('foi: „educație” a doua oară primește sufixul (2)', 'Educație' in nume and 'educație (2)' in nume, nume)
    verifica('foi: caracterele interzise din nume sunt înlocuite', not any(c in n for n in nume for c in '/\\?*[]:'), nume)
    ws = wb['Rezumat']
    verifica('foi: link în interiorul fișierului către foaia „Educație” (celula A1)', ws['A5'].hyperlink is not None and ws['A5'].hyperlink.location == "'Educație'!A1", ws['A5'].hyperlink and ws['A5'].hyperlink.location)
    # foaia cu nume prea lung se numește altfel decât a fost cerută: linkul trebuie să arate spre numele final
    final = [n for n in nume if n.startswith('Sănătate')][0]
    verifica('foi: linkul către o foaie cu nume tăiat folosește numele final', ws['A6'].hyperlink is not None and ws['A6'].hyperlink.location == "'%s'!A1" % final, ws['A6'].hyperlink and ws['A6'].hyperlink.location)
    verifica('foi: linkul către o foaie care nu există nu se creează', ws['A9'].hyperlink is None, ws['A9'].hyperlink)
    ws = wb['Educație']
    verifica('foi: hyperlinkul extern din foaia de sector rămâne extern', ws['A5'].hyperlink is not None and ws['A5'].hyperlink.target == 'https://amp.gov.md/A0', ws['A5'].hyperlink and ws['A5'].hyperlink.target)
    ws = wb[final]
    verifica('foi: totalul fiecărei foi are formula lui, pe zona proprie', ws['C7'].value == '=SUBTOTAL(109,C5:C6)', ws['C7'].value)
    with zipfile.ZipFile(cale) as z:
        parti = z.namelist()
        sheets = {i: z.read('xl/worksheets/sheet%d.xml' % i).decode('utf-8') for i in range(1, 7)}
        wbx = z.read('xl/workbook.xml').decode('utf-8')
        ct = z.read('[Content_Types].xml').decode('utf-8')
        rel1 = z.read('xl/worksheets/_rels/sheet1.xml.rels').decode('utf-8')
        rel2 = z.read('xl/worksheets/_rels/sheet2.xml.rels').decode('utf-8')
        rel4 = z.read('xl/worksheets/_rels/sheet4.xml.rels').decode('utf-8')
        ch = {k: z.read('xl/charts/chart%d.xml' % k).decode('utf-8') for k in (1, 2, 3) if 'xl/charts/chart%d.xml' % k in parti}
    verifica('foi: graficele sunt numerotate pe tot fișierul (chart1..3, drawing1..3), fără dubluri',
             all(('xl/charts/chart%d.xml' % k) in parti and ('xl/drawings/drawing%d.xml' % k) in parti and ('xl/drawings/_rels/drawing%d.xml.rels' % k) in parti for k in (1, 2, 3))
             and 'xl/charts/chart4.xml' not in parti, [p for p in parti if 'chart' in p or 'drawing' in p])
    verifica('foi: fiecare foaie cu grafic îl leagă de desenul ei', 'drawing1.xml' in rel1 and 'drawing2.xml' in rel2 and 'drawing3.xml' in rel4, (rel1, rel2, rel4))
    verifica('foi: graficele arată spre foile lor', "'Rezumat'!" in ch.get(1, '') and "'Toate proiectele'!" in ch.get(2, '') and ("'%s'!" % final) in ch.get(3, ''), {k: re.findall(r'<c:f>(.*?)</c:f>', v)[:2] for k, v in ch.items()})
    verifica('foi: tipurile de conținut le numesc pe toate (6 foi, 3 grafice)', all(('/xl/worksheets/sheet%d.xml' % i) in ct for i in range(1, 7)) and ct.count('drawingml.chart+xml') == 3 and ct.count('drawing+xml') == 3, ct.count('chart+xml'))
    verifica('foi: doar prima foaie e selectată', 'tabSelected="1"' in sheets[1] and not any('tabSelected' in sheets[i] for i in range(2, 7)))
    verifica('foi: filtrul și rândurile care se repetă au localSheetId potrivit fiecărei foi',
             all(('localSheetId="%d"' % i) in wbx for i in range(6)) and wbx.count('_xlnm._FilterDatabase') == 6 and wbx.count('_xlnm.Print_Titles') == 6, wbx[wbx.find('<definedNames>'):][:400])
    verifica('foi: eticheta foii poate avea culoare proprie', 'tabColor rgb="FF2E7D32"' in sheets[2] and 'tabColor rgb="FF1E4B8C"' in sheets[1], re.findall(r'tabColor rgb="(\w+)"', sheets[2]))
    verifica('foi: subsolul e pus pe fiecare foaie', all('nistor.vivi.md' in sheets[i] for i in range(1, 7)))
    with zipfile.ZipFile(cale) as z:
        verifica('foi: stilurile și șirurile sunt comune (un singur styles.xml și un singur sharedStrings.xml)', [p for p in z.namelist() if p.endswith('styles.xml') or p.endswith('sharedStrings.xml')] == ['xl/styles.xml', 'xl/sharedStrings.xml'])
    # o singură foaie, cum se făcea înainte: aceeași ieșire în privința numelui și a primei foi
    GEN_UNA = r"""
const X = require(process.argv[2]); const fs = require('fs');
(async () => { const b = await X.construieste({ limba: 'en', foaie: 'Sheet one', titlu: 'T', subtitlu: 's', antete: ['A', 'B'], tipuri: 'tn', randuri: [['x', 1]] });
  fs.writeFileSync(process.argv[3] + '/una.xlsx', Buffer.from(await b.arrayBuffer())); })();
"""
    (tmp / 'gen_una.js').write_text(GEN_UNA, encoding='utf-8')
    subprocess.run(['node', str(tmp / 'gen_una.js'), str(AICI / 'xlsx-stil.js'), str(tmp)], check=True, timeout=60)
    w1 = openpyxl.load_workbook(tmp / 'una.xlsx')
    verifica('foi: specificația cu o singură foaie merge ca înainte', w1.sheetnames == ['Sheet one'] and w1.worksheets[0]['A5'].value == 'x')

    # ---- tabele late, ca registrul acordurilor: coloane înghețate, bandă de grupe, culori pe etichete, înălțime limitată
    GEN_REG = r"""
const X = require(process.argv[2]); const fs = require('fs');
const salveaza = async (nume, spec) => { const b = await X.construieste(spec); fs.writeFileSync(process.argv[3] + '/' + nume + '.xlsx', Buffer.from(await b.arrayBuffer())); };
(async () => {
  const lung = 'Acord cu un nume foarte lung care se întinde pe mai multe rânduri '.repeat(5);
  const mediu = 'Un nume mediu, care încape pe două rânduri la lățimea coloanei, dar nu pe trei';
  await salveaza('reg', {
    limba: 'ro', foaie: 'Reg', titlu: 'Titlu registru', subtitlu: 'Subtitlu registru', subsol: 'x',
    antete: ['Acord', 'Tip', 'Da', 'Val', 'Curs', 'Fragment', 'Suma'], tipuri: 'ttmcrle',
    randuri: [['Scurt', 'Grant', 'da', 'EUR', 1.1383, 'text lung ' + 'x'.repeat(300), 10], [mediu, 'Împrumut', '', 'USD', 0.83998, 'f', 20],
              [lung, 'Altceva', 'da', 'MDL', 1, 'f', 30], ['Cu "ghilimele"', 'Cu "ghilimele"', '', 'EUR', 1, 'f', 40]],
    latimi: { 0: 40, 1: 14, 5: 30 }, inghetaColoane: 1, maxLinii: 3, paginiLatime: 'auto', total: { eticheta: 'Total', suma: [6] },
    grupe: [{ de: 0, pana: 1, text: 'ACORD' }, { de: 2, pana: 3, text: 'MARCAJE' }, { de: 4, pana: 6, text: 'CIFRE' }],
    etichete: { 1: [['Grant', 'FF155A52', 'FFE0F2F0'], ['Cu "ghilimele"', 'FFA3311E', 'FFFBE8E3']], 2: [['da', 'FF16794F', 'FFE3F4EC']] }
  });
  const nume = []; for (let i = 0; i < 12; i++) nume.push('C' + i);
  const lat = {}; nume.forEach((_, i) => { lat[i] = 20; });
  await salveaza('lat', { limba: 'ro', foaie: 'Lat', titlu: 'Titlu pe un tabel lat', subtitlu: 'Context pe un tabel lat', antete: nume, tipuri: 'tttttttttttt',
    randuri: [nume.map((_, i) => 'v' + i)], latimi: lat, paginiLatime: 'auto', note: ['O notă pe un tabel lat.'] });
  await salveaza('grupe_rele', { limba: 'ro', foaie: 'Grupe', titlu: 'T', subtitlu: 's', antete: ['A', 'B', 'C', 'D'], tipuri: 'tttt', randuri: [['1', '2', '3', '4']],
    grupe: [{ de: 2, pana: 3, text: 'DOUA' }, { de: 0, pana: 1, text: 'UNU' }, { de: 1, pana: 2, text: 'SUPRAPUS' }, { de: 5, pana: 99, text: 'AFARA' }, { de: 3, pana: 1, text: 'INVERS' }] });
})().catch(e => { console.error('EROARE', e && e.stack || e); process.exit(1); });
"""
    (tmp / 'gen_reg.js').write_text(GEN_REG, encoding='utf-8')
    r = subprocess.run(['node', str(tmp / 'gen_reg.js'), str(AICI / 'xlsx-stil.js'), str(tmp)], capture_output=True, text=True, timeout=120)
    verifica('registru: Node generează fișierele de probă pentru tabele late', r.returncode == 0, (r.stderr or r.stdout)[-300:])
    for nm in ('reg', 'lat', 'grupe_rele'):
        verifica(f'registru: {nm}: arhivă ZIP validă, cu toate părțile XML bine formate', (tmp / (nm + '.xlsx')).exists() and not parti_xml_valide(tmp / (nm + '.xlsx')), parti_xml_valide(tmp / (nm + '.xlsx')))
    wb = openpyxl.load_workbook(tmp / 'reg.xlsx'); ws = wb.worksheets[0]
    verifica('registru: prima coloană și antetul rămân pe loc la derulare (pane înghețat în ambele sensuri)', ws.freeze_panes == 'B6', ws.freeze_panes)
    with zipfile.ZipFile(tmp / 'reg.xlsx') as z:
        foaie = z.read('xl/worksheets/sheet1.xml').decode('utf-8')
        wbx = z.read('xl/workbook.xml').decode('utf-8')
    verifica('registru: panourile înghețate au selecția fiecărui panou (topRight, bottomLeft, bottomRight)',
             'xSplit="1"' in foaie and 'ySplit="5"' in foaie and 'activePane="bottomRight"' in foaie and all(('<selection pane="%s"' % p) in foaie for p in ('topRight', 'bottomLeft', 'bottomRight')), re.findall(r'<pane [^>]*>', foaie))
    verifica('registru: antetul tabelului e sub banda de grupe (rândul 5), cu filtre doar pe tabel', ws['A5'].value == 'Acord' and ws.auto_filter.ref == 'A5:G9', (ws['A5'].value, ws.auto_filter.ref))
    verifica('registru: banda de grupe are denumirile pe rândul 4, unite peste coloanele lor',
             [ws['A4'].value, ws['C4'].value, ws['E4'].value] == ['ACORD', 'MARCAJE', 'CIFRE'] and {'A4:B4', 'C4:D4', 'E4:G4'} <= {str(m) for m in ws.merged_cells.ranges},
             ([ws['A4'].value, ws['C4'].value, ws['E4'].value], sorted(str(m) for m in ws.merged_cells.ranges)))
    verifica('registru: grupele suprapuse, inversate sau în afara tabelului sunt ignorate',
             (lambda w: [w['A4'].value, w['B4'].value, w['C4'].value, w['D4'].value] == ['UNU', None, 'DOUA', None] and w['A5'].value == 'A')(openpyxl.load_workbook(tmp / 'grupe_rele.xlsx').worksheets[0]),
             [openpyxl.load_workbook(tmp / 'grupe_rele.xlsx').worksheets[0].cell(row=4, column=c).value for c in range(1, 5)])
    # tipurile noi: m (marcaj), c (centrat), r (curs), l (text lung pe un rând)
    verifica('registru: tipul „m" — marcaj centrat, îngroșat, verde', ws['C6'].alignment.horizontal == 'center' and ws['C6'].font.b and ws['C6'].font.color.rgb == 'FF16794F' and ws['C6'].value == 'da',
             (ws['C6'].alignment.horizontal, ws['C6'].font.b, ws['C6'].font.color.rgb))
    verifica('registru: tipul „c" — text scurt centrat, fără îngroșare', ws['D6'].alignment.horizontal == 'center' and not ws['D6'].font.b and ws['D6'].value == 'EUR', (ws['D6'].alignment.horizontal, ws['D6'].font.b))
    verifica('registru: tipul „r" — cursul are patru zecimale și rămâne număr', ws['E6'].number_format == '#,##0.0000' and abs(ws['E6'].value - 1.1383) < 1e-9 and abs(ws['E7'].value - 0.83998) < 1e-9, (ws['E6'].number_format, ws['E6'].value))
    verifica('registru: tipul „l" — text lung pe un singur rând (fără încadrare), „t" — cu încadrare', ws['F6'].alignment.wrap_text is not True and ws['A6'].alignment.wrap_text is True, (ws['F6'].alignment.wrap_text, ws['A6'].alignment.wrap_text))
    verifica('registru: antetul tipurilor de text e la stânga, al celorlalte centrat', ws['A5'].alignment.horizontal == 'left' and ws['F5'].alignment.horizontal == 'left' and ws['C5'].alignment.horizontal == 'center' and ws['E5'].alignment.horizontal == 'center',
             [ws[c + '5'].alignment.horizontal for c in 'ACEF'])
    # înălțimea rândurilor
    verifica('registru: rândul care încape în 3 linii nu primește înălțime fixă (o potrivește aplicația), cel cu text lung e tăiat la 3 linii (45 pt)',
             ws.row_dimensions[6].height is None and ws.row_dimensions[7].height is None and ws.row_dimensions[8].height == 45, [ws.row_dimensions[r].height for r in (6, 7, 8)])
    verifica('registru: coloana „l” cu 310 de caractere nu înalță rândul', ws.row_dimensions[6].height is None)
    verifica('registru: cu înălțime limitată, celulele de date încep sus (se taie jos, nu pe mijloc)', ws['A8'].alignment.vertical == 'top' and ws['G8'].alignment.vertical == 'top', (ws['A8'].alignment.vertical, ws['G8'].alignment.vertical))
    verifica('registru: fără înălțime limitată, celulele rămân centrate pe verticală (ca la exporturile din dashboard)',
             openpyxl.load_workbook(tmp / 'top.xlsx').worksheets[0]['A5'].alignment.vertical == 'center' if (tmp / 'top.xlsx').exists() else True)
    # culori pe etichete
    reguli = {}
    for rng in ws.conditional_formatting:
        for r_ in rng.rules:
            reguli.setdefault(str(rng.sqref), []).append(r_)
    de_tip = reguli.get('B6:B9', [])
    verifica('registru: două reguli de culoare pe coloana „Tip", pe valori exacte (cellIs / equal)', len(de_tip) == 2 and all(r_.type == 'cellIs' and r_.operator == 'equal' for r_ in de_tip), [(r_.type, r_.operator) for r_ in de_tip])
    verifica('registru: ghilimelele din etichete sunt dublate în formulă („Cu ""ghilimele""")', [r_.formula[0] for r_ in de_tip] == ['"Grant"', '"Cu ""ghilimele"""'], [r_.formula for r_ in de_tip])
    d0 = de_tip[0].dxf if de_tip else None
    verifica('registru: eticheta „Grant” are textul verde-închis îngroșat și fundalul verzui', d0 is not None and d0.font.b and d0.font.color.rgb == 'FF155A52' and d0.fill.bgColor.rgb == 'FFE0F2F0', d0 and (d0.font.color.rgb, d0.fill.bgColor.rgb))
    zebra = [r_ for r_s in reguli.values() for r_ in r_s if r_.type == 'expression']
    verifica('registru: culorile etichetelor au prioritate peste dungile alternante (număr de prioritate mai mic)', zebra and de_tip and max(r_.priority for r_ in de_tip) < zebra[0].priority, ([r_.priority for r_ in de_tip], [r_.priority for r_ in zebra]))
    verifica('registru: prioritățile regulilor sunt unice', len([r_.priority for r_s in reguli.values() for r_ in r_s]) == len({r_.priority for r_s in reguli.values() for r_ in r_s}))
    # tipărirea
    verifica('registru: la tipărire se repetă prima coloană și antetul, iar tabelul larg încape pe mai multe pagini în lățime',
             ws.print_title_cols == '$A:$A' and ws.print_title_rows == '$5:$5' and ws.page_setup.fitToWidth >= 1, (ws.print_title_cols, ws.print_title_rows, ws.page_setup.fitToWidth))
    verifica('registru: totalul adună doar coloana cerută, cu formulă SUBTOTAL', ws['G10'].value == '=SUBTOTAL(109,G6:G9)' and ws['A10'].value == 'Total', (ws['A10'].value, ws['G10'].value))
    # titlul pe tabele late: unit doar peste coloanele care încap pe ecran, fundalul pe toată lățimea
    wl = openpyxl.load_workbook(tmp / 'lat.xlsx').worksheets[0]
    unite = sorted(str(m) for m in wl.merged_cells.ranges)
    verifica('lat: titlul, contextul și nota sunt unite doar peste primele 8 coloane (160 de caractere), nu peste toate 12', {'A1:H1', 'A2:H2', 'A7:H7'} <= set(unite) and 'A1:L1' not in unite and 'A7:L7' not in unite, unite)
    verifica('lat: fundalul benzii de titlu continuă pe toată lățimea (ultima coloană are aceeași culoare)', wl['L1'].fill.fgColor.rgb == wl['A1'].fill.fgColor.rgb == 'FF14233C', (wl['L1'].fill.fgColor.rgb, wl['A1'].fill.fgColor.rgb))
    verifica('lat: tabelul de 240 de caractere se tipărește pe 2 pagini în lățime, nu pe una', wl.page_setup.fitToWidth == 2, wl.page_setup.fitToWidth)
    verifica('lat: fără coloane înghețate, repetarea la tipărire e doar a antetului', wl.print_title_cols is None and wl.print_title_rows == '$4:$4', (wl.print_title_cols, wl.print_title_rows))
    verifica('lat: panoul înghețat rămâne doar pe rânduri (A5), ca la exporturile obișnuite', wl.freeze_panes == 'A5', wl.freeze_panes)
    wt = openpyxl.load_workbook(tmp / 'top.xlsx').worksheets[0]
    verifica('top: la un tabel îngust titlul rămâne unit pe toată lățimea (A1:D1) și tipărirea e pe o pagină',
             'A1:D1' in {str(m) for m in wt.merged_cells.ranges} and wt.page_setup.fitToWidth == 1, (sorted(str(m) for m in wt.merged_cells.ranges)[:3], wt.page_setup.fitToWidth))
    if '--pastreaza' in sys.argv:
        shutil.copy(cale, AICI.parent / 'foi_probă.xlsx')
finally:
    shutil.rmtree(tmp, ignore_errors=True)


# ============================================================ 2. în browser
if '--browser' in sys.argv:
    print()
    print('Dashboardul, în browser: ⚙ → „Descarcă Excel” pe panouri')
    sys.path.insert(0, str(AICI))
    import legatura_acorduri as lga
    from playwright.sync_api import sync_playwright

    srv = lga._Server(('127.0.0.1', 0), functools.partial(lga._Liniste, directory=str(AICI)))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    baza = 'http://127.0.0.1:%d/' % srv.server_address[1]

    # activități IATI de probă, ca să existe ce exporta din paginile IATI și D-portal fără rețea
    JS_IATI = '''async () => {
      const P = PROJECTS.filter(p => p.id && p.ang > 0).slice(0, 300);
      const acts = [];
      P.forEach((p, i) => { if(i % 3) return;
        acts.push({aid: 'XM-DAC-MOCK-' + i, title: p.denumire, reporting: 'World Bank', reporting_ref: '44000', status_code: 2 + (i % 3),
          day_start: Math.floor(Date.UTC(2010 + (i % 12), 0, 1) / 86400000), day_end: null, commitment_eur: Math.round(p.ang * 0.9),
          spend_eur: i % 4 ? Math.round(p.deb) : Math.round(p.deb * 0.5), description: '', sectors: [{code: '11110', name: 'Education policy'}]}); });
      IATI_ACTIVITIES = acts; IATI_FETCH_DATE = new Date().toISOString().slice(0, 10); IATI_LOAD_ERROR = null; IATI_STALE = false;
      await iatiApply();
      return acts.length;
    }'''
    JS_REG = '''([vid, i]) => {
      const p = document.querySelectorAll('#' + vid + ' .panel')[i];
      let d = p.__dateCsv; if(typeof d === 'function') d = d();
      const h = p.querySelector('.panel-head h2'), sp = h.querySelector('[data-i18n]');
      return {rows: d ? d.rows.length : -1, cols: d ? d.cols.length : -1, opt: d && d.opt ? {grafic: !!d.opt.grafic, total: !!d.opt.total, randTotal: !!d.opt.randTotal, totalFinal: !!d.opt.totalFinal} : null,
              titlu: (sp || h).textContent.replace(/\\s+/g, ' ').trim()};
    }'''
    JS_INDEX = '''([vid, cheie]) => [...document.querySelectorAll('#' + vid + ' .panel')].findIndex(p => {
      const h = p.querySelector('.panel-head h2'); const s = h && h.querySelector('[data-i18n]');
      return p.id === cheie || (s && s.getAttribute('data-i18n') === cheie); })'''

    DESCARCARI = Path(tempfile.mkdtemp(prefix='teste_xlsx_dl_'))

    def salveaza_dl(d, nume):                      # Playwright dă fișierul fără extensie: îl salvăm sub un nume cu extensie
        cale = DESCARCARI / nume
        d.save_as(str(cale))
        return str(cale)

    ERORI = []
    try:
        with sync_playwright() as p:
            b = p.chromium.launch()
            for limba in ('ro', 'en'):
                ctx = b.new_context(viewport={'width': 1500, 'height': 1000}, accept_downloads=True)
                pg = ctx.new_page()
                pg.on('pageerror', lambda e: ERORI.append(str(e)[:200]))
                pg.route('**/*', lambda r: r.continue_() if r.request.url.startswith(baza) else r.abort())
                pg.goto(baza + 'index.html', wait_until='networkidle', timeout=120000)
                pg.wait_for_timeout(3500)
                pg.evaluate('window.finishIntro && finishIntro()')
                if limba == 'en':
                    pg.evaluate("setLanguage('en')"); pg.wait_for_timeout(1500)
                pg.evaluate(JS_IATI); pg.wait_for_timeout(1000)

                def descarca_xlsx(vid, cheie, actiune='xlsx'):
                    # evenimentele „scroll” ajung cu o întârziere și pot închide meniul: a doua încercare e legitimă
                    for incercare in (1, 2):
                        try:
                            return _descarca_xlsx(vid, cheie, actiune)
                        except Exception:
                            if incercare == 2: raise
                            try: pg.keyboard.press('Escape')
                            except Exception: pass
                            pg.wait_for_timeout(900)

                def _descarca_xlsx(vid, cheie, actiune='xlsx'):
                    i = pg.evaluate(JS_INDEX, [vid, cheie])
                    if i < 0:
                        return None, None, None
                    reg = pg.evaluate(JS_REG, [vid, i])
                    panel = pg.locator('#' + vid + ' .panel').nth(i)
                    panel.scroll_into_view_if_needed(); pg.wait_for_timeout(450)
                    roata = panel.locator(':scope > .panel-head .roata-panou').first
                    roata.scroll_into_view_if_needed(); pg.wait_for_timeout(450)
                    roata.click()
                    pg.wait_for_selector('.meniu-panou.deschis', timeout=4000)
                    iteme = pg.eval_on_selector_all('.meniu-panou.deschis [data-actiune]', 'els => els.map(e => e.dataset.actiune)')
                    with pg.expect_download(timeout=40000) as dl:
                        pg.locator('.meniu-panou.deschis [data-actiune="%s"]' % actiune).click(timeout=8000)
                    return dl.value, reg, iteme

                print(f'  [{limba}]')
                tinte = [('mainView', 'topDonorsTitle', True, True), ('mainView', 'sectorsTitle', True, True), ('mainView', 'donorTableTitle', True, False),
                         ('mainView', 'projectsListTitle', True, False), ('advancedView', 'advLineTitle', True, True),
                         ('advancedView', 'advExecTitle', True, True), ('advancedView', 'advHeatmapTitle', True, False), ('advancedView', 'advRegionTitle', True, True),
                         ('advancedView', 'advTyTitle', True, True), ('iatiView', 'iatiDifTitle', False, False), ('iatiView', 'iatiResultTitle', False, False),
                         ('dportalView', 'dpSectorTitle', True, True), ('dportalView', 'dpProjectsTitle', True, False), ('legaturiView', 'legListTitle', False, False)]
                verifica(f'{limba} · „Analize avansate": panoul „Comparație între perioade" a fost scos (nici în pagină, nici funcția lui)',
                         pg.evaluate("() => !document.getElementById('cmpPanel') && !document.getElementById('cmpTitlu') && typeof window.renderComparatie === 'undefined'"))
                vedere_curenta = 'mainView'
                for vid, cheie, cu_total, cu_grafic in tinte:
                    if vid != vedere_curenta:
                        pg.evaluate({'advancedView': 'showAdvancedView()', 'iatiView': 'showIatiView()', 'dportalView': 'showDportalView()', 'legaturiView': 'showLegaturiView()'}[vid])
                        pg.wait_for_timeout(4500 if vid == 'legaturiView' else 3000)
                        vedere_curenta = vid
                    try:
                        d, reg, iteme = descarca_xlsx(vid, cheie)
                    except Exception as e:
                        verifica(f'{limba} · {cheie}: descărcarea prin ⚙', False, str(e).split('\n')[0][:150]); continue
                    if d is None:
                        verifica(f'{limba} · {cheie}: panoul există', False); continue
                    nume = d.suggested_filename
                    cale = salveaza_dl(d, '%s-%s.xlsx' % (limba, cheie))
                    ws = openpyxl.load_workbook(cale).worksheets[0]
                    z = zipfile.ZipFile(cale)
                    o = reg['opt'] or {}
                    verifica(f'{limba} · {cheie}: meniul are 3 acțiuni (imagine, Excel, CSV)', iteme == ['copiaza', 'xlsx', 'csv'], iteme)
                    verifica(f'{limba} · {cheie}: fișier .xlsx cu data în nume ({nume})', nume.endswith('.xlsx') and re.search(r'\d{4}-\d{2}-\d{2}', nume) is not None, nume)
                    verifica(f'{limba} · {cheie}: titlul din A1 e al panoului, subtitlul are sursa și data',
                             ws['A1'].value.split(' · ')[0] == reg['titlu'].split(' · ')[0] and re.search(r'(Sursa|Source): ', ws['A2'].value or '') and re.search(r'(Generat|Generated): ', ws['A2'].value or ''),
                             (ws['A1'].value, reg['titlu'], (ws['A2'].value or '')[:80]))
                    # filtrul acoperă exact rândurile de date (antetul e pe rândul 4); când ultimul rând al panoului e totalul (totalFinal),
                    # acesta iese din filtru și devine rândul de total de sub tabel
                    n_date = reg['rows'] - (1 if o.get('totalFinal') else 0)
                    verifica(f'{limba} · {cheie}: toate cele {reg["rows"]} rânduri sunt în fișier, cu filtru și antet înghețat',
                             reg['rows'] > 0 and (ws.auto_filter.ref or '').endswith(str(4 + n_date)) and ws.freeze_panes == 'A5'
                             and ws.cell(row=4 + n_date, column=1).value not in (None, ''), (reg['rows'], ws.auto_filter.ref, ws.freeze_panes))
                    if o.get('total') or o.get('randTotal') or o.get('totalFinal'):
                        et = ws.cell(row=5 + n_date, column=1).value
                        verifica(f'{limba} · {cheie}: rând de total sub date', isinstance(et, str) and et.lower().startswith('total'), et)
                    verifica(f'{limba} · {cheie}: grafic Excel doar unde panoul are grafic', ('xl/charts/chart1.xml' in z.namelist()) == bool(o.get('grafic') and reg['rows'] > 1),
                             (o, 'xl/charts/chart1.xml' in z.namelist()))
                    verifica(f'{limba} · {cheie}: fișierul e valid (arhivă + XML bine format)', not parti_xml_valide(cale), parti_xml_valide(cale))

                # CSV-ul a rămas, ca a doua opțiune
                d, reg, iteme = descarca_xlsx('legaturiView' if vedere_curenta == 'legaturiView' else vedere_curenta, 'legListTitle', 'csv') if vedere_curenta == 'legaturiView' else (None, None, None)
                if d is not None:
                    brut = Path(salveaza_dl(d, '%s-date-brute.csv' % limba)).read_bytes()
                    verifica(f'{limba} · CSV (date brute) rămas în meniu: BOM UTF-8 și separator „;”', d.suggested_filename.endswith('.csv') and brut[:3] == b'\xef\xbb\xbf' and b';' in brut[:400], d.suggested_filename)

                # butoanele de sub tabele
                pg.evaluate('showLegaturiView()'); pg.wait_for_timeout(4000)
                with pg.expect_download(timeout=40000) as dl:
                    pg.click('#legCsvBtn')
                ws = openpyxl.load_workbook(salveaza_dl(dl.value, '%s-legaturi.xlsx' % limba)).worksheets[0]
                verifica(f'{limba} · „De la lege la bani”: butonul de sub filtre dă Excel cu 17 coloane și linkuri spre registru', dl.value.suggested_filename.endswith('.xlsx') and ws.max_column == 17
                         and sum(1 for row in ws.iter_rows(min_row=5, min_col=17, max_col=17) if row[0].hyperlink) > 100, (dl.value.suggested_filename, ws.max_column))
                pg.evaluate('showIatiView()'); pg.wait_for_timeout(3000)
                with pg.expect_download(timeout=40000) as dl:
                    pg.click('#iatiDiffCsvBtn')
                ws = openpyxl.load_workbook(salveaza_dl(dl.value, '%s-iati-diferente.xlsx' % limba)).worksheets[0]
                verifica(f'{limba} · IATI: „Descarcă diferențele” dă Excel cu 19 coloane, doar proiectele potrivite, sortate după diferență',
                         dl.value.suggested_filename.endswith('.xlsx') and ws.max_column == 19 and ws['K5'].value is not None
                         and abs(ws['K5'].value) >= abs(ws['K6'].value or 0), (dl.value.suggested_filename, ws.max_column, ws['K5'].value, ws['K6'].value))
                ctx.close()
            b.close()
    finally:
        srv.shutdown()
        shutil.rmtree(DESCARCARI, ignore_errors=True)
    verifica('paginile nu au dat erori JavaScript în timpul exporturilor', not ERORI, ERORI[:3])

print()
if picat:
    print(f'Au picat {len(picat)} teste.'); sys.exit(1)
print('Toate testele exportului Excel au trecut.')
